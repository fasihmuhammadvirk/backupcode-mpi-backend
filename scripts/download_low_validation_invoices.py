"""
Download PDFs and create Excel reports for invoices with validation_score < 95.

Organizes output by vendor folder under analysis/low_validation_invoices/.
"""

import json
import os
import sys
from collections import defaultdict
from pathlib import Path

from dotenv import load_dotenv

# Load .env from project root
project_root = Path(__file__).resolve().parent.parent
load_dotenv(dotenv_path=project_root / ".env")

import boto3
import psycopg2
import psycopg2.extras
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
OUTPUT_DIR = project_root / "analysis" / "low_validation_invoices"

DB_CONFIG = {
    "host": os.getenv("POSTGRES_HOST"),
    "port": os.getenv("POSTGRES_PORT"),
    "dbname": os.getenv("POSTGRES_DB"),
    "user": os.getenv("POSTGRES_USERNAME"),
    "password": os.getenv("POSTGRES_PASSWORD"),
}

S3_BUCKET = os.getenv("INVOICES_PDFS_BUCKET_NAME", "invoice-website-mpi")
S3_REGION = os.getenv("INVOICES_PDFS_REGION", "us-east-2")
S3_ACCESS_KEY = os.getenv("INVOICES_PDFS_AWS_SERVER_PUBLIC_KEY")
S3_SECRET_KEY = os.getenv("INVOICES_PDFS_AWS_SERVER_SECRET_KEY")

QUERY = """
WITH latest AS (
  SELECT med.id, med.s3_file_key, med.validation_score, med.extracted_at,
    med.pdf_extraction_queue_id, med.status, med.confidence_score, med.model_name,
    med.raw_data,
    ROW_NUMBER() OVER (PARTITION BY med.s3_file_key ORDER BY med.extracted_at DESC) as rn
  FROM raw.model_extracted_data med
  WHERE med.deleted_at IS NULL
)
SELECT l.id, l.s3_file_key, l.validation_score, l.extracted_at, l.status,
  l.confidence_score, l.model_name, l.raw_data,
  CASE WHEN peq.status = 'pending' THEN true ELSE false END as is_pending
FROM latest l
LEFT JOIN raw.pdf_extraction_queue peq ON l.pdf_extraction_queue_id = peq.id
WHERE l.rn = 1
  AND l.validation_score < 95
ORDER BY l.validation_score ASC;
"""

# Excel column definitions for per-vendor sheets
DETAIL_COLUMNS = [
    "id",
    "s3_file_key",
    "validation_score",
    "status",
    "confidence_score",
    "model_name",
    "extracted_at",
    "is_pending",
    "ref_no",
    "date",
    "vendor",
    "po_no",
    "customer_job",
    "buyer",
    "sub_total",
    "tax",
    "freight",
    "line_item_count",
    "has_header_details",
    "has_validation_failed_checks",
    "validation_failed_checks",
    "raw_data",
]

SUMMARY_COLUMNS = [
    "vendor",
    "invoice_count",
    "pending_count",
    "not_pending_count",
    "avg_validation_score",
    "min_validation_score",
    "max_validation_score",
]

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def extract_vendor(s3_file_key: str) -> str:
    """Return the second segment of the s3 key path (the vendor folder)."""
    parts = s3_file_key.split("/")
    if len(parts) >= 2:
        return parts[1]
    return "unknown_vendor"


def safe_get(raw_data: dict | None, *keys, default=None):
    """Safely traverse nested dict keys."""
    if raw_data is None:
        return default
    val = raw_data
    for k in keys:
        if isinstance(val, dict):
            val = val.get(k)
        else:
            return default
    return val if val is not None else default


def parse_row(row: dict) -> dict:
    """Enrich a DB row with fields extracted from raw_data."""
    raw = row.get("raw_data")
    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except (json.JSONDecodeError, TypeError):
            raw = {}
    if raw is None:
        raw = {}

    line_items = raw.get("line_items")
    line_item_count = len(line_items) if isinstance(line_items, list) else 0

    has_header = "_header_details" in raw if isinstance(raw, dict) else False

    # Validation failed checks -- could be in raw_data or top-level
    failed_checks = safe_get(raw, "validation_failed_checks", default=None)
    has_failed = bool(failed_checks)
    if isinstance(failed_checks, list):
        failed_checks_str = "; ".join(str(c) for c in failed_checks)
    elif failed_checks is not None:
        failed_checks_str = str(failed_checks)
    else:
        failed_checks_str = ""

    return {
        "id": str(row["id"]),
        "s3_file_key": row["s3_file_key"],
        "validation_score": row["validation_score"],
        "status": row["status"],
        "confidence_score": row["confidence_score"],
        "model_name": row["model_name"],
        "extracted_at": str(row["extracted_at"]) if row["extracted_at"] else "",
        "is_pending": row["is_pending"],
        "ref_no": safe_get(raw, "ref_no", default=""),
        "date": safe_get(raw, "date", default=""),
        "vendor": safe_get(raw, "vendor", default=""),
        "po_no": safe_get(raw, "po_no", default=""),
        "customer_job": safe_get(raw, "customer_job", default=""),
        "buyer": safe_get(raw, "buyer", default=""),
        "sub_total": safe_get(raw, "sub_total", default=""),
        "tax": safe_get(raw, "tax", default=""),
        "freight": safe_get(raw, "freight", default=""),
        "line_item_count": line_item_count,
        "has_header_details": has_header,
        "has_validation_failed_checks": has_failed,
        "validation_failed_checks": failed_checks_str,
        "raw_data": json.dumps(raw, default=str),
    }


def style_header(ws):
    """Apply bold font and fill to the first row."""
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def auto_width(ws):
    """Auto-size column widths (capped at 60)."""
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                val_len = len(str(cell.value)) if cell.value else 0
                if val_len > max_len:
                    max_len = val_len
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


def write_vendor_excel(vendor: str, rows: list[dict], vendor_dir: Path):
    """Write a per-vendor Excel workbook."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Extraction Data"
    ws.append(DETAIL_COLUMNS)
    for r in rows:
        ws.append([r[c] for c in DETAIL_COLUMNS])
    style_header(ws)
    auto_width(ws)
    path = vendor_dir / f"{vendor}_extraction_data.xlsx"
    wb.save(str(path))
    return path


def write_summary_excel(vendor_stats: dict, output_dir: Path):
    """Write summary.xlsx with per-vendor aggregations."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(SUMMARY_COLUMNS)
    for vendor in sorted(vendor_stats.keys()):
        s = vendor_stats[vendor]
        ws.append([
            vendor,
            s["count"],
            s["pending"],
            s["not_pending"],
            round(s["total_score"] / s["count"], 2) if s["count"] else 0,
            s["min_score"],
            s["max_score"],
        ])
    style_header(ws)
    auto_width(ws)
    path = output_dir / "summary.xlsx"
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():
    # 1. Query the database
    print("[1/4] Querying database for low-validation invoices...")
    conn = psycopg2.connect(**DB_CONFIG)
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute(QUERY)
            rows = cur.fetchall()
    finally:
        conn.close()

    print(f"       Found {len(rows)} invoices")
    if not rows:
        print("       No invoices found. Exiting.")
        sys.exit(0)

    # 2. Parse rows and group by vendor
    print("[2/4] Parsing extraction data and grouping by vendor...")
    vendor_rows: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        parsed = parse_row(row)
        vendor = extract_vendor(row["s3_file_key"])
        parsed["_vendor_folder"] = vendor
        parsed["_s3_key"] = row["s3_file_key"]
        vendor_rows[vendor].append(parsed)

    print(f"       {len(vendor_rows)} vendors identified")

    # 3. Create vendor dirs and download PDFs
    print("[3/4] Downloading PDFs from S3...")
    s3_client = boto3.client(
        "s3",
        region_name=S3_REGION,
        aws_access_key_id=S3_ACCESS_KEY,
        aws_secret_access_key=S3_SECRET_KEY,
    )

    downloaded = 0
    failed = 0
    for vendor, inv_rows in sorted(vendor_rows.items()):
        vendor_dir = OUTPUT_DIR / vendor
        vendor_dir.mkdir(parents=True, exist_ok=True)

        for r in inv_rows:
            s3_key = r["_s3_key"]
            filename = s3_key.split("/")[-1]
            local_path = vendor_dir / filename

            if local_path.exists():
                downloaded += 1
                continue

            try:
                s3_client.download_file(S3_BUCKET, s3_key, str(local_path))
                downloaded += 1
            except Exception as e:
                print(f"       WARN: Failed to download {s3_key}: {e}")
                failed += 1

    print(f"       Downloaded {downloaded} PDFs ({failed} failures)")

    # 4. Create Excel files
    print("[4/4] Creating Excel reports...")
    vendor_stats: dict[str, dict] = {}

    for vendor, inv_rows in sorted(vendor_rows.items()):
        vendor_dir = OUTPUT_DIR / vendor

        # Per-vendor Excel
        write_vendor_excel(vendor, inv_rows, vendor_dir)

        # Accumulate stats for summary
        scores = [r["validation_score"] for r in inv_rows if r["validation_score"] is not None]
        pending_count = sum(1 for r in inv_rows if r["is_pending"])
        vendor_stats[vendor] = {
            "count": len(inv_rows),
            "pending": pending_count,
            "not_pending": len(inv_rows) - pending_count,
            "total_score": sum(scores),
            "min_score": min(scores) if scores else 0,
            "max_score": max(scores) if scores else 0,
        }

    # Summary Excel
    summary_path = write_summary_excel(vendor_stats, OUTPUT_DIR)
    print(f"       Summary written to {summary_path}")
    print("[OK] Done.")


if __name__ == "__main__":
    main()
